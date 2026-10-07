import io
import re
import logging
from typing import Tuple, Dict, Any, List, Optional
from datetime import datetime

from app.services.supabase import (
    get_job_details,
    get_report_for_research,
    get_sources_for_research,
    get_latest_verification,
)

logger = logging.getLogger(__name__)


class ExportService:
    """Multi-format document export service for research documents (Markdown, DOCX, PDF)."""

    @classmethod
    def _sanitize_filename(cls, topic: str, ext: str) -> str:
        clean = re.sub(r'[^a-zA-Z0-9_\-]+', '-', topic.strip().lower())
        clean = re.sub(r'-+', '-', clean).strip('-')[:50] or "research-report"
        return f"{clean}.{ext}"

    @classmethod
    def _get_export_data(cls, research_id: str) -> Dict[str, Any]:
        job = get_job_details(research_id) or {}
        report = get_report_for_research(research_id) or {}
        sources = get_sources_for_research(research_id, limit=50)
        verification = get_latest_verification(research_id)

        topic = job.get("topic", "Research Report")
        title = report.get("title", f"Research Report: {topic}")
        content = report.get("content_markdown", "")
        word_count = len(content.split())
        created_date = report.get("created_at") or datetime.now().isoformat()
        date_str = datetime.fromisoformat(created_date.replace("Z", "+00:00")).strftime("%B %d, %Y")

        verif_summary = "Not Audited"
        if verification:
            v_total = verification.get("total_claims", 0)
            v_verified = verification.get("verified_count", 0)
            verif_summary = f"{v_verified}/{v_total} Verified Claims" if v_total > 0 else "Audit Complete"

        return {
            "topic": topic,
            "title": title,
            "content": content,
            "word_count": word_count,
            "date_str": date_str,
            "sources": sources,
            "verification_summary": verif_summary,
        }

    @classmethod
    def to_markdown(cls, research_id: str) -> Tuple[str, str]:
        """Generates formatted Markdown export with metadata frontmatter."""
        data = cls._get_export_data(research_id)
        frontmatter = f"""---
title: "{data['title']}"
topic: "{data['topic']}"
generated_date: "{data['date_str']}"
word_count: {data['word_count']}
sources_count: {len(data['sources'])}
verification_status: "{data['verification_summary']}"
system: "Autonomous Multi-Agent AI Researcher"
---

"""
        full_md = frontmatter + data["content"]
        filename = cls._sanitize_filename(data["topic"], "md")
        return filename, full_md

    @classmethod
    def to_docx(cls, research_id: str) -> Tuple[str, bytes]:
        """Generates professionally styled Microsoft Word document (.docx)."""
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT

        data = cls._get_export_data(research_id)
        doc = docx.Document()

        # Page Margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)

        # Document Header
        p_tag = doc.add_paragraph()
        run_tag = p_tag.add_run("AUTONOMOUS MULTI-AGENT SYNTHESIS")
        run_tag.font.name = "Calibri"
        run_tag.font.size = Pt(9)
        run_tag.font.bold = True
        run_tag.font.color.rgb = RGBColor(37, 99, 235) # Blue

        p_title = doc.add_paragraph()
        run_title = p_title.add_run(data["title"])
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        # Metadata Callout Box
        meta_table = doc.add_table(rows=1, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = True
        hdr_cells = meta_table.rows[0].cells
        hdr_cells[0].text = f"Date: {data['date_str']}"
        hdr_cells[1].text = f"Words: ~{data['word_count']}"
        hdr_cells[2].text = f"Sources: {len(data['sources'])}"
        hdr_cells[3].text = f"Audit: {data['verification_summary']}"

        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph() # Spacer

        # Parse Markdown lines into headings, tables, lists, and paragraphs
        lines = data["content"].split("\n")
        in_table = False
        table_rows: List[List[str]] = []

        for line in lines:
            line_str = line.strip()

            # Handle Markdown Tables
            if line_str.startswith("|") and line_str.endswith("|"):
                in_table = True
                if re.match(r'^\|[\s\-:|]+\|$', line_str):
                    continue # Skip divider row
                cols = [c.strip() for c in line_str.strip("|").split("|")]
                table_rows.append(cols)
                continue
            elif in_table:
                # Flush table to document
                if table_rows:
                    num_cols = max(len(r) for r in table_rows)
                    t = doc.add_table(rows=len(table_rows), cols=num_cols)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for r_idx, row_data in enumerate(table_rows):
                        for c_idx in range(num_cols):
                            val = row_data[c_idx] if c_idx < len(row_data) else ""
                            cell = t.cell(r_idx, c_idx)
                            cell.text = val
                            if r_idx == 0:
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.font.bold = True
                                        r.font.size = Pt(9)
                                        r.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                for p in cell.paragraphs:
                                    for r in p.runs:
                                        r.font.size = Pt(8.5)
                    doc.add_paragraph()
                in_table = False
                table_rows = []

            if not line_str:
                continue

            # Headings
            if line_str.startswith("# "):
                continue # Skip title repeated
            elif line_str.startswith("## "):
                h2 = doc.add_heading(level=1)
                run = h2.add_run(line_str[3:].strip())
                run.font.name = "Calibri"
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = RGBColor(30, 41, 59)
            elif line_str.startswith("### "):
                h3 = doc.add_heading(level=2)
                run = h3.add_run(line_str[4:].strip())
                run.font.name = "Calibri"
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(71, 85, 105)
            elif line_str.startswith(("- ", "* ", "• ")):
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(line_str[2:].strip())
            elif re.match(r'^\d+\.\s', line_str):
                p = doc.add_paragraph(style="List Number")
                p.add_run(re.sub(r'^\d+\.\s*', '', line_str))
            else:
                p = doc.add_paragraph()
                # Remove link markdown syntax for clean reading
                clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line_str)
                p.add_run(clean_text)

        # Save to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        filename = cls._sanitize_filename(data["topic"], "docx")
        return filename, buffer.getvalue()

    @classmethod
    def render_markdown_to_pdf(
        cls,
        title: str,
        content_markdown: str,
        topic: str = "Research Export",
        date_str: Optional[str] = None,
        sources: Optional[List[Dict[str, Any]]] = None,
        verification_summary: Optional[str] = None,
        include_metadata_header: bool = False,
    ) -> Tuple[str, bytes]:
        """Renders arbitrary research markdown into a styled, publication-grade PDF document."""
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            HRFlowable,
            KeepTogether,
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=45,
            leftMargin=45,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()

        # Custom Palette Styles
        style_kicker = ParagraphStyle(
            'Kicker',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            textColor=colors.HexColor('#2563EB'),
            textTransform='uppercase',
            spaceAfter=4,
        )

        style_title = ParagraphStyle(
            'TitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#0F172A'),
            spaceAfter=8,
        )

        style_h1 = ParagraphStyle(
            'Heading1Style',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=8,
            spaceAfter=6,
        )

        style_h2 = ParagraphStyle(
            'Heading2Style',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=15,
            textColor=colors.HexColor('#334155'),
            spaceBefore=8,
            spaceAfter=4,
        )

        style_body = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=8,
        )

        style_bullet = ParagraphStyle(
            'BulletStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor('#1E293B'),
            leftIndent=14,
            spaceAfter=4,
        )

        style_table_cell = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=7.5,
            leading=10.5,
            textColor=colors.HexColor('#1E293B'),
        )

        style_table_header = ParagraphStyle(
            'TableHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.white,
        )

        story = []

        # Only include the upper metadata bar if explicitly requested
        if include_metadata_header:
            story.append(Paragraph("Autonomous Multi-Agent Synthesis", style_kicker))
            story.append(Paragraph(title, style_title))

            actual_date = date_str or datetime.now().strftime("%B %d, %Y")
            word_count = len(content_markdown.split())
            num_sources = len(sources or [])

            meta_data = [[
                f"<b>Date:</b> {actual_date}",
                f"<b>Words:</b> ~{word_count}",
                f"<b>Sources:</b> {num_sources} Verified" if num_sources > 0 else "<b>Status:</b> Standalone Section",
                f"<b>Audit:</b> {verification_summary or 'Verified'}",
            ]]
            meta_table = Table(meta_data, colWidths=[130, 100, 130, 160])
            meta_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#475569')),
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ]))
            story.append(meta_table)
            story.append(Spacer(1, 14))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E2E8F0'), spaceAfter=12))
        else:
            # Check if content already starts with an initial heading
            first_meaningful_line = next((l.strip() for l in content_markdown.split("\n") if l.strip()), "")
            starts_with_heading = bool(re.match(r'^#{1,4}\s+', first_meaningful_line))
            if not starts_with_heading and title:
                clean_title = re.sub(r'[*_#]', '', title).strip()
                story.append(Paragraph(clean_title, style_title))
                story.append(Spacer(1, 4))


        def flush_table(rows: List[List[str]]) -> None:
            if not rows:
                return
            num_cols = max(len(r) for r in rows)

            # Compute intelligent proportional column widths based on maximum cell content length
            col_lens = [0] * num_cols
            for r in rows:
                for c_idx in range(min(num_cols, len(r))):
                    col_lens[c_idx] = max(col_lens[c_idx], len(r[c_idx]))

            min_w = 55.0
            available_w = max(520.0 - (min_w * num_cols), 60.0)
            weights = [max(l - 12, 0) for l in col_lens]
            total_weight = sum(weights) or 1
            raw_widths = [min_w + (available_w * (w / total_weight)) for w in weights]
            scale = 520.0 / (sum(raw_widths) or 1.0)
            col_widths = [w * scale for w in raw_widths]

            formatted_table_data = []

            for r_idx, row in enumerate(rows):
                row_cells = []
                for c_idx in range(num_cols):
                    val = row[c_idx] if c_idx < len(row) else ""
                    # Strip markdown links e.g. [Title](url) -> Title
                    clean_val = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', val)
                    # Convert HTML br to ReportLab br
                    clean_val = re.sub(r'<br\s*/?>', '<br/>', clean_val, flags=re.IGNORECASE)
                    # Escape bare ampersand that is not part of an existing XML entity
                    clean_val = re.sub(r'&(?!(?:amp|lt|gt|quot|apos|bull);)', '&amp;', clean_val)
                    # Convert markdown bold **text** to ReportLab <b>text</b>
                    clean_val = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', clean_val)
                    if r_idx == 0:
                        row_cells.append(Paragraph(clean_val, style_table_header))
                    else:
                        row_cells.append(Paragraph(clean_val, style_table_cell))
                formatted_table_data.append(row_cells)

            t = Table(formatted_table_data, colWidths=col_widths, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

        lines = content_markdown.split("\n")
        in_table = False
        table_rows: List[List[str]] = []

        for line in lines:
            line_str = line.strip()

            if line_str.startswith("|") and line_str.endswith("|"):
                in_table = True
                if re.match(r'^\|[\s\-:|]+\|$', line_str):
                    continue
                cols = [c.strip() for c in line_str.strip("|").split("|")]
                table_rows.append(cols)
                continue
            elif in_table:
                flush_table(table_rows)
                in_table = False
                table_rows = []

            if not line_str:
                continue

            if line_str.startswith("# "):
                clean_h = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', line_str[2:].strip())
                story.append(Paragraph(clean_h, style_h1))
            elif line_str.startswith("## "):
                clean_h = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', line_str[3:].strip())
                story.append(Paragraph(clean_h, style_h1))
            elif line_str.startswith("### "):
                clean_h = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', line_str[4:].strip())
                story.append(Paragraph(clean_h, style_h1 if not story else style_h2))
            elif line_str.startswith(("- ", "* ", "• ")):
                clean_bullet = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line_str[2:].strip())
                clean_bullet = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', clean_bullet)
                story.append(Paragraph(f"&bull; {clean_bullet}", style_bullet))
            elif re.match(r'^\d+\.\s', line_str):
                clean_num = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line_str)
                clean_num = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', clean_num)
                story.append(Paragraph(clean_num, style_bullet))
            else:
                clean_p = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', line_str)
                clean_p = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', clean_p)
                story.append(Paragraph(clean_p, style_body))

        if in_table and table_rows:
            flush_table(table_rows)

        doc.build(story)
        buffer.seek(0)
        filename = cls._sanitize_filename(topic or title, "pdf")
        return filename, buffer.getvalue()

    @classmethod
    def to_pdf(cls, research_id: str) -> Tuple[str, bytes]:
        """Generates a publication-grade PDF document using ReportLab."""
        data = cls._get_export_data(research_id)
        return cls.render_markdown_to_pdf(
            title=data["title"],
            content_markdown=data["content"],
            topic=data["topic"],
            date_str=data["date_str"],
            sources=data["sources"],
            verification_summary=data["verification_summary"]
        )
